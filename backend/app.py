from flask import Flask, request, jsonify, abort, send_from_directory
from flask_sqlalchemy import SQLAlchemy
import openai
import os
import docx
import PyPDF2
import subprocess
import uuid
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from flask_cors import CORS
from datetime import datetime

# Laden der Umgebungsvariablen aus der .env-Datei
load_dotenv()

# Flask-Konfiguration
app = Flask(__name__)
CORS(app)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///db.sqlite'  # SQLite-Datenbank
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = 'uploads'
basedir = os.path.abspath(os.path.dirname(__file__))
app.config['PRESENTATION_FOLDER'] = os.path.join(basedir, 'presentations')
app.config['CUSTOMTHEME_FOLDER'] = os.path.join(basedir, 'customtheme')
app.config['TEMPLATE_FOLDER'] = os.path.join(basedir, 'templates')
app.config['ALLOWED_EXTENSIONS'] = {'txt', 'pdf', 'docx'}

# Datenbank initialisieren
db = SQLAlchemy(app)

# OpenAI API-Schlüssel
openai.api_key = os.getenv('OPENAI_API_KEY')

# Definition der Datenbankmodelle
class TextEntry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_input = db.Column(db.Text, nullable=False)
    gpt_response = db.Column(db.Text, nullable=False)
    segmentation = db.Column(db.Text, nullable=True)
    key_points = db.Column(db.Text, nullable=True)
    is_temporary = db.Column(db.Boolean, default=False)

class Presentation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    summary_id = db.Column(db.Integer)
    format = db.Column(db.String(10))
    markdown_content = db.Column(db.Text)
    file_path = db.Column(db.String(200))
    title = db.Column(db.String(200))
    is_temporary = db.Column(db.Boolean, default=False)

# Hilfsfunktionen für die Text- und Dateiverarbeitung
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def extract_text_from_file(file):
    filename = secure_filename(file.filename)
    file_extension = filename.rsplit('.', 1)[1].lower()

    if file_extension == 'txt':
        return file.read().decode('utf-8')
    elif file_extension == 'pdf':
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in range(len(pdf_reader.pages)):
            text += pdf_reader.pages[page].extract_text()
        return text
    elif file_extension == 'docx':
        doc = docx.Document(file)
        return '\n'.join([paragraph.text for paragraph in doc.paragraphs])
    else:
        return None

@app.route('/api/process-text', methods=['POST'])
def process_text():
    texts = []
    response_data = []

    # Verarbeiten von hochgeladenen Dateien
    for key in request.files:
        file = request.files[key]
        if file and allowed_file(file.filename):
            text = extract_text_from_file(file)
            if text:
                texts.append(text)

    # Falls kein Text aus Dateien vorhanden ist, verwende den eingegebenen Text
    if not texts:
        text = request.form.get('text', '')
        if text:
            texts.append(text)

    if not texts:
        return jsonify({'error': 'Kein gültiger Text bereitgestellt'}), 400

    run_segmentation = request.form.get('run_segmentation') == 'on'
    run_summary = request.form.get('run_summary') == 'on'
    run_key_points = request.form.get('run_key_points') == 'on'
    is_temporary_str = request.form.get('is_temporary', 'false')
    is_temporary = is_temporary_str.lower() == 'true'

    try:
        summary_length = int(request.form.get('summary_length', 250))
    except ValueError:
        summary_length = 250

    # Verarbeitung für jeden Text
    for text in texts:
        single_response = {}
        if run_segmentation:
            segmentation_prompt = (
                "Bitte analysiere den folgenden Text und identifiziere alle Überschriften "
                "sowie die entsprechenden Textsegmente. Gib die Überschriften und die "
                "zugehörigen Segmente strukturiert wieder.\n\n"
                f"Text:\n{text}"
            )
            segmentation_response = openai.ChatCompletion.create(
                model="gpt-4o-mini",
                temperature=0.3,
                messages=[
                    {"role": "system", "content": "Du bist ein Experte für Textanalyse und antwortest auf Deutsch."},
                    {"role": "user", "content": segmentation_prompt}
                ]
            )
            segmentation = segmentation_response.choices[0].message['content'].strip()
            single_response['segmentation'] = segmentation

        if run_summary:
            summary_prompt = (
                f"Schreibe eine detaillierte Zusammenfassung des folgenden Textes. "
                f"Die Zusammenfassung soll möglichst nah an {summary_length} Wörtern sein, "
                f"und sie sollte so viele relevante Details wie möglich enthalten, um den Inhalt umfassend abzudecken.\n\nText:\n{text}"
            )
            summary_response = openai.ChatCompletion.create(
                model="gpt-4o-mini",
                temperature=0.3,
                messages=[
                    {"role": "system", "content": "Du bist ein Experte im Zusammenfassen von Texten und im Extrahieren von Schlüsselpunkten und antwortest auf Deutsch."},
                    {"role": "user", "content": summary_prompt}
                ]
            )
            summary = summary_response.choices[0].message['content'].strip()
            single_response['summary'] = summary

        if run_key_points:
            key_points_prompt = f"Extrahiere die wichtigsten Schlüsselpunkte aus dem folgenden Text:\n\n{text}"
            key_points_response = openai.ChatCompletion.create(
                model="gpt-4o-mini",
                temperature=0.3,
                messages=[
                    {"role": "system", "content": "Du bist ein Experte für Schlüsselpunkt-Extraktion und antwortest auf Deutsch."},
                    {"role": "user", "content": key_points_prompt}
                ]
            )
            key_points = key_points_response.choices[0].message['content'].strip()
            single_response['key_points'] = key_points

        if not is_temporary:
            new_entry = TextEntry(
                user_input=text,
                gpt_response=single_response.get('summary', ''),
                segmentation=single_response.get('segmentation', ''),
                key_points=single_response.get('key_points', ''),
                is_temporary=is_temporary
            )
            db.session.add(new_entry)
            db.session.commit()
            single_response['id'] = new_entry.id

        response_data.append(single_response)

    return jsonify(response_data)

# Route zum Abrufen der Zusammenfassungen
@app.route('/api/summaries', methods=['GET'])
def get_summaries():
    summaries = TextEntry.query.all()
    summaries_list = [
        {
            'id': entry.id,
            'user_input': entry.user_input,
            'gpt_response': entry.gpt_response,
            'segmentation': entry.segmentation,
            'key_points': entry.key_points
        }
        for entry in summaries
    ]
    return jsonify(summaries_list)

# Route zum Abrufen oder Löschen einer Zusammenfassung
@app.route('/api/summaries/<int:id>', methods=['GET', 'DELETE'])
def get_or_delete_summary(id):
    if request.method == 'GET':
        summary = db.session.get(TextEntry, id)
        if summary:
            return jsonify({
                'id': summary.id,
                'user_input': summary.user_input,
                'gpt_response': summary.gpt_response,
                'segmentation': summary.segmentation,
                'key_points': summary.key_points
            }), 200
        else:
            return jsonify({'error': 'Zusammenfassung nicht gefunden'}), 404
    elif request.method == 'DELETE':
        summary = db.session.get(TextEntry, id)
        if summary:
            db.session.delete(summary)
            db.session.commit()
            return jsonify({"message": "Zusammenfassung erfolgreich gelöscht"}), 200
        else:
            return jsonify({"error": "Zusammenfassung nicht gefunden"}), 404

@app.route('/api/create-presentation', methods=['POST'])
def create_presentation():
    # Daten aus dem Formular abrufen
    summary_id = request.form.get('summary_id')
    raw_text = request.form.get('summary_text')  # Direkt übergebener Text
    format = request.form.get('format', '').lower()
    document_theme = request.form.get('theme', None)
    document_colortheme = request.form.get('colortheme', None)
    template = request.form.get('template', 'none')
    is_temporary_str = request.form.get('is_temporary', 'false')
    is_temporary = is_temporary_str.lower() == 'true'  # Inkognito-Modus

    # Hochgeladenes benutzerdefiniertes Template abrufen
    custom_template_file = request.files.get('custom_template')

    # Überprüfen, ob entweder summary_id oder raw_text bereitgestellt wird
    if summary_id:
        # Zusammenfassung aus der Datenbank abrufen
        summary = db.session.get(TextEntry, summary_id)
        if not summary:
            return jsonify({'error': 'Zusammenfassung nicht gefunden'}), 404
        user_text = summary.user_input
    elif raw_text:
        # Verwende direkt den übergebenen Text
        user_text = raw_text
    else:
        return jsonify({'error': 'Weder summary_id noch Text vorhanden'}), 400

    # Extrahieren von Titel und Autor aus dem Text
    title_author_prompt = (
        "Analysiere den folgenden wissenschaftlichen Text und identifiziere den Titel und den Autor des Dokuments. "
        "Gib die Informationen im folgenden Format zurück:\n\n"
        "Titel: [Titel des Dokuments]\n"
        "Autor: [Name des Autors]\n\n"
        f"Text:\n{user_text}"
    )

    try:
        title_author_response = openai.ChatCompletion.create(
            model="gpt-4o-mini",
            temperature=0.3,
            messages=[
                {"role": "system", "content": "Du bist ein Experte für Textanalyse und antwortest auf Deutsch."},
                {"role": "user", "content": title_author_prompt}
            ]
        )
        title_author_content = title_author_response.choices[0].message['content'].strip()

        # Verarbeiten von Titel und Autor
        title_line = [line for line in title_author_content.splitlines() if line.startswith("Titel:")]
        author_line = [line for line in title_author_content.splitlines() if line.startswith("Autor:")]

        document_title = title_line[0].replace("Titel:", "").strip() if title_line else "Kein Titel erkannt"
        document_author = author_line[0].replace("Autor:", "").strip() if author_line else "Kein Autor erkannt"

    except Exception as e:
        print(f"Fehler beim Abrufen des Titels und Autors: {e}")
        document_title = "Kein Titel erkannt"
        document_author = "Kein Autor erkannt"

    # Generierung des Markdown-Inhalts
    markdown_prompt = (
        f"Erstelle eine wissenschaftliche Präsentation im Markdown-Format basierend auf dem folgenden wissenschaftlichen Text. "
        "Gib das Ergebnis direkt als reinen Markdown-Text zurück, ohne es in Codeblöcke oder Anführungszeichen zu setzen.\n\n"
        f"Text:\n{user_text}\n\n"
        "Die Präsentation muss genau 10 Folien umfassen. Jede Folie darf nur 3 Aufzählungspunkte enthalten. "
        "Die erste Folie muss eine Agenda sein und sollte nicht zu den 10 Folien zählen. "
        "Verwende keine Zahlen oder Buchstaben zur Nummerierung der Überschriften und Aufzählungspunkte. "
        "Erstelle klare, einfache Aufzählungspunkte ohne zusätzliche Markierungen oder Sonderzeichen. "
        "Beginne jede neue Folie mit einer Überschrift der Ebene 2 (`## Titel der Folie`). "
        "Nutze für die Überschriften die Markdown-Formatierung (## für Folientitel). "
        "Verwende keine doppelten Folientitel oder Untertitel."
    )

    try:
        response = openai.ChatCompletion.create(
            model="gpt-4o-mini",
            temperature=0.3,
            messages=[
                {"role": "system", "content": "Du bist ein Experte für wissenschaftliche Präsentationen und antwortest auf Deutsch."},
                {"role": "user", "content": markdown_prompt}
            ]
        )
        markdown_content = response.choices[0].message['content'].strip()

        # Entferne eventuell vorhandene Codeblock-Markierungen
        if markdown_content.startswith('```') and markdown_content.endswith('```'):
            lines = markdown_content.split('\n')
            if lines[0].startswith('```'):
                lines = lines[1:]
            if lines[-1].startswith('```'):
                lines = lines[:-1]
            markdown_content = '\n'.join(lines)

        print("Markdown-Inhalt erfolgreich von GPT erhalten.")
    except Exception as e:
        print(f"Fehler beim Abrufen der GPT-Antwort: {e}")
        return jsonify({'error': 'Fehler beim Abrufen der GPT-Antwort'}), 500

    # Füge den YAML-Header für Metadaten hinzu (Titel, Autor und Theme)
    metadata_header = f"---\ntitle: '{document_title}'\nauthor: '{document_author}'\n"

    if document_theme and document_theme != 'default':
        metadata_header += f"theme: '{document_theme}'\n"

    if document_colortheme and document_colortheme != 'default':
        metadata_header += f"colortheme: '{document_colortheme}'\n"

    metadata_header += "---\n\n"

    markdown_content_with_metadata = metadata_header + markdown_content

    # Verarbeitung des benutzerdefinierten Templates
    template_path = None
    if template == 'custom' and custom_template_file:
        # Speichern der hochgeladenen Datei
        filename = secure_filename(custom_template_file.filename)
        template_directory = os.path.join(app.root_path, 'templates', 'custom')
        if not os.path.exists(template_directory):
            os.makedirs(template_directory)
        template_path = os.path.join(template_directory, filename)
        custom_template_file.save(template_path)
        print(f"Benutzerdefiniertes Template gespeichert unter: {template_path}")
    elif template == 'medieninformatik':
        # Verwenden des vordefinierten Medieninformatik-Templates
        template_path = os.path.join(app.root_path, 'templates', 'medieninformatik.pptx')
        print(f"Medieninformatik-Template verwendet: {template_path}")
    else:
        # Kein Template oder Standard-Template
        template_path = None
        print("Kein spezielles Template ausgewählt.")

    # Generierung einer eindeutigen ID für die Präsentation
    if is_temporary or not summary_id:
        presentation_id = str(uuid.uuid4())
    else:
        presentation_id = str(summary_id)

    # Generierung der Präsentation
    try:
        output_file_path, success = generate_presentation(
            presentation_id=presentation_id,
            format=format,
            markdown_content=markdown_content_with_metadata,
            template=template,
            template_path=template_path
        )
    except Exception as e:
        print(f"Fehler bei der Generierung der Präsentation: {e}")
        return jsonify({'error': 'Fehler bei der Erstellung der Präsentation'}), 500

    if success:
        if not is_temporary:
            # Präsentation dauerhaft speichern
            new_presentation = Presentation(
                summary_id=summary_id,
                format=format,
                markdown_content=markdown_content_with_metadata,
                file_path=output_file_path,
                title=document_title,
                is_temporary=False  # Dauerhafte Speicherung
            )
            db.session.add(new_presentation)
            db.session.commit()
            return jsonify({
                'id': new_presentation.id,
                'title': new_presentation.title,
                'format': new_presentation.format,
                'is_temporary': new_presentation.is_temporary
            })
        else:
            # Präsentationsdaten zurückgeben
            return jsonify({
                'id': presentation_id,
                'title': document_title,
                'format': format,
                'is_temporary': True
            })
    else:
        return jsonify({'error': 'Fehler bei der Erstellung der Präsentation'}), 500

def generate_presentation(presentation_id, format, markdown_content, template, template_path=None):
    basedir = os.path.abspath(os.path.dirname(__file__))
    template_folder = os.path.join(basedir, 'templates')
    presentation_folder = os.path.join(basedir, 'presentations')

    # Stelle sicher, dass der Präsentationsordner existiert
    if not os.path.exists(presentation_folder):
        os.makedirs(presentation_folder)

    # Erstelle den Markdown-Dateinamen
    markdown_filename = f"presentation_{presentation_id}.md"
    markdown_file_path = os.path.join(presentation_folder, markdown_filename)

    with open(markdown_file_path, 'w', encoding='utf-8') as md_file:
        md_file.write(markdown_content)

    # Setze die richtige Dateiendung basierend auf dem Format
    if format == "beamer":
        output_filename = f"presentation_{presentation_id}.pdf"  # Beamer ergibt eine PDF
    elif format == "latex":
        output_filename = f"presentation_{presentation_id}.tex"
    elif format == "pptx":
        output_filename = f"presentation_{presentation_id}.pptx"
    else:
        output_filename = f"presentation_{presentation_id}.{format}"

    output_file_path = os.path.join(presentation_folder, output_filename)

    # Erstelle den Pandoc-Befehl
    pandoc_command = ["pandoc", markdown_file_path, "--output", output_file_path]

    if format == "beamer":
        pandoc_command.extend(["--to=beamer", "--pdf-engine=xelatex"])
    elif format == "latex":
        pandoc_command.extend(["--to=beamer", "--standalone"])
    elif format == "pdf":
        pandoc_command.extend(["--pdf-engine=xelatex"])
    elif format == "pptx":
        pandoc_command.extend(["-t", "pptx", "--standalone"])

        # Template-Validierung
        if template == 'medieninformatik':
            default_template_path = os.path.join(template_folder, 'Medieninformatik-Template.pptx')
            if os.path.exists(default_template_path):
                pandoc_command.extend(["--reference-doc", default_template_path])
                print(f"Medieninformatik-Template wird verwendet: {default_template_path}")
            else:
                print(f"Medieninformatik-Template nicht gefunden. Standardvorlage wird verwendet.")
        elif template == 'custom' and template_path:
            if os.path.exists(template_path):
                pandoc_command.extend(["--reference-doc", template_path])
                print(f"Benutzerdefiniertes Template wird verwendet: {template_path}")
            else:
                print(f"Benutzerdefiniertes Template nicht gefunden. Standardvorlage wird verwendet.")

    # Fallback für unbekannte Formate
    if format not in ["beamer", "latex", "pdf", "pptx"]:
        print(f"Unbekanntes Format '{format}'. Es wird 'pdf' verwendet.")
        pandoc_command.extend(["--pdf-engine=xelatex"])

    try:
        subprocess.run(pandoc_command, check=True)
        print(f"Präsentation wurde erfolgreich im Format '{format}' generiert.")

        # Fügen Sie einen Zeitstempel hinzu (optional)
        current_time = datetime.now().timestamp()
        os.utime(output_file_path, (current_time, current_time))

        return output_file_path, True
    except subprocess.CalledProcessError as e:
        print(f"Fehler bei der Pandoc-Konvertierung: {e}")
        return None, False

@app.route('/api/presentations', methods=['GET'])
def get_presentations():
    presentations = Presentation.query.all()
    presentations_data = [
        {
            'id': presentation.id,
            'summary_id': presentation.summary_id,
            'title': presentation.title,
            'format': presentation.format,
            'file_path': presentation.file_path,
            'markdown_content': presentation.markdown_content  # Hier wird der markdown_content abgerufen
        }
        for presentation in presentations
    ]
    return jsonify(presentations_data)

@app.route('/api/presentations/<int:presentation_id>/delete', methods=['DELETE'])
def delete_presentation(presentation_id):
    print(f"Request to delete presentation with ID: {presentation_id}")
    
    presentation = db.session.get(Presentation, presentation_id)
    
    if not presentation:
        print(f"Presentation with ID {presentation_id} not found")
        return jsonify({'error': 'Präsentation nicht gefunden'}), 404

    try:
        # Lösche die Präsentation aus der Datenbank
        db.session.delete(presentation)
        db.session.commit()

        # Lösche die Präsentationsdatei von der Festplatte
        if presentation.file_path and os.path.exists(presentation.file_path):
            os.remove(presentation.file_path)
            print(f"Deleted file at path: {presentation.file_path}")
        else:
            print(f"File not found at path: {presentation.file_path}")
        
        # Lösche die zugehörige Markdown-Datei
        markdown_filename = f"presentation_{presentation_id}.md"
        markdown_file_path = os.path.join(app.config['PRESENTATION_FOLDER'], markdown_filename)
        
        if os.path.exists(markdown_file_path):
            os.remove(markdown_file_path)
            print(f"Deleted markdown file at path: {markdown_file_path}")
        else:
            print(f"Markdown file not found at path: {markdown_file_path}")
        
        return jsonify({'message': 'Präsentation erfolgreich gelöscht'})
    
    except Exception as e:
        print(f"Error occurred while deleting presentation: {e}")
        return jsonify({'error': 'Fehler beim Löschen der Präsentation'}), 500

@app.route('/presentations/<string:presentation_id>/download', methods=['GET'])
def download_presentation(presentation_id):
    print(f"Download function called for presentation ID: {presentation_id}")

    # Validierung der presentation_id
    is_uuid = False
    try:
        uuid_obj = uuid.UUID(presentation_id, version=4)
        is_uuid = True
    except ValueError:
        # Die ID ist keine gültige UUID, prüfen ob es eine numerische ID ist
        if not presentation_id.isdigit():
            print("Invalid presentation ID")
            return jsonify({'error': 'Ungültige Präsentations-ID'}), 400

    # Wenn es sich um eine numerische ID handelt, casten wir sie zu int
    if not is_uuid:
        presentation_id_int = int(presentation_id)
    else:
        presentation_id_int = None  # Für UUID-basierte IDs

    presentation = None
    if not is_uuid:
        # Suche die Präsentation in der Datenbank
        presentation = db.session.get(Presentation, presentation_id_int)

    if presentation:
        # Präsentation in der Datenbank gefunden
        file_path = presentation.file_path
        is_temporary = presentation.is_temporary
        summary_id = presentation.summary_id
        # Hole die zugehörige Zusammenfassung, falls vorhanden
        summary = db.session.get(TextEntry, summary_id) if summary_id else None
    else:
        # Präsentation nicht in der Datenbank gefunden, versuchen wir, die Datei direkt zu finden
        presentation_folder = app.config.get('PRESENTATION_FOLDER', os.path.join(app.root_path, 'presentations'))
        possible_formats = ['pdf', 'pptx', 'latex']
        file_path = None
        is_temporary = True  # Annahme, dass Präsentationen außerhalb der DB temporär sind
        for fmt in possible_formats:
            temp_file_path = os.path.join(presentation_folder, f"presentation_{presentation_id}.{fmt}")
            if os.path.exists(temp_file_path):
                file_path = temp_file_path
                summary = None
                break
        if not file_path:
            print("Presentation not found")
            return jsonify({'error': 'Präsentation nicht gefunden'}), 404

    # Überprüfen, ob die Datei existiert
    if not os.path.exists(file_path):
        print("File does not exist")
        return abort(404)

    # Debugging-Ausgabe
    print(f"Attempting to download file: {file_path}")

    # Bereite den Download vor
    directory = os.path.dirname(file_path)
    filename = os.path.basename(file_path)
    response = send_from_directory(directory, filename, as_attachment=True)

    # Prüfen, ob die Präsentation als temporär gekennzeichnet ist
    if is_temporary:
        try:
            # Lösche die Präsentation aus der Datenbank
            if presentation:
                db.session.delete(presentation)
                if summary and summary.is_temporary:
                    db.session.delete(summary)
                db.session.commit()

            # Lösche die Präsentationsdatei
            if os.path.exists(file_path):
                os.remove(file_path)
                print(f"Deleted presentation file: {file_path}")

            # Lösche die zugehörige .md-Datei
            markdown_file = os.path.join(app.config['PRESENTATION_FOLDER'], f"presentation_{presentation_id}.md")
            if os.path.exists(markdown_file):
                os.remove(markdown_file)
                print(f"Deleted markdown file: {markdown_file}")
        except Exception as e:
            print(f"Fehler beim Löschen der temporären Datei: {e}")

    return response

def is_valid_uuid(uuid_to_test, version=4):
    try:
        uuid_obj = uuid.UUID(uuid_to_test, version=version)
    except ValueError:
        return False
    return str(uuid_obj) == uuid_to_test

# Hauptprogramm
if __name__ == '__main__':
    if not os.path.exists(app.config['UPLOAD_FOLDER']):
        os.makedirs(app.config['UPLOAD_FOLDER'])
    if not os.path.exists(app.config['PRESENTATION_FOLDER']):
        os.makedirs(app.config['PRESENTATION_FOLDER'])
    if not os.path.exists(app.config['CUSTOMTHEME_FOLDER']):
        os.makedirs(app.config['CUSTOMTHEME_FOLDER'])
    if not os.path.exists(app.config['TEMPLATE_FOLDER']):
        os.makedirs(app.config['TEMPLATE_FOLDER'])
    
    with app.app_context():
        db.create_all()

    app.run(debug=True)
